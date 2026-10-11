#!/usr/bin/env python3
"""Calcule l'affectation qui réduit le pire trajet, puis l'écrit dans Supabase.

Un magasin peut rester vide. S'il reçoit des étudiants, il reste dans sa capacité.
Chaque zone DIF en reçoit au moins 20.
Les durées viennent du cache Google, arrivée un mardi à 9 h. Le cache Navitia n'est pas lu.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import subprocess
import tempfile
import threading
import time
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urlencode
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

warnings.filterwarnings("ignore", message="urllib3 v2 only supports OpenSSL")

import pulp
import requests
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent
CACHE_PATH = ROOT / "cache" / "temps_trajet.json"
CAR_CACHE_PATH = ROOT / "cache" / "temps_voiture.json"
GOOGLE_CACHE_PATH = ROOT / "cache" / "temps_google.json"
GOOGLE_ADDRESSES_PATH = ROOT.parent / "adresses_google.xlsx"
GOOGLE_RULES = "arrivee-9h-google-transit"
NAVITIA_URL = "https://prim.iledefrance-mobilites.fr/marketplace/v2/navitia/journeys"
OSRM_URL = "https://router.project-osrm.org/route/v1/driving"
ZONES = ("DIF01", "DIF02", "DIF03", "DIF04")
IDF_DEPTS = {"75", "77", "78", "91", "92", "93", "94", "95"}
ZONE_MIN = 20
OUTSIDE_RANK = 99
NAVITIA_PER_SECOND = 8
WALKING_MPS = 6 / 3.6
TRIP_RULES = "arrivee-9h-marche-6-sans-court"
UNREACHABLE = 10**9
RANK_STEPS = (99, 10, 9, 8, 7, 6, 5, 4, 3, 2)
# numéro de vote, code stocké, famille, objectif de temps
ALGORITHMS = (
    (5, "adresses_pire", "adresses", "pire"),
)
# L'étudiant est affecté à l'un de ces magasins, dans chaque répartition.
FORCED_STORES = {
    "victoire.basse@hec.edu": ("Ormesson",),
    "andre.vieira@hec.edu": ("Ormesson",),
}
_SOLVER = None


def make_solver():
    """CBC quand le binaire tourne, sinon HiGHS (Apple Silicon)."""
    global _SOLVER
    if _SOLVER is not None:
        return _SOLVER
    try:
        cbc = pulp.PULP_CBC_CMD(msg=False, timeLimit=180)
        if cbc.available():
            probe = pulp.LpProblem("probe", pulp.LpMinimize)
            variable = pulp.LpVariable("probe_var", cat="Binary")
            probe += variable
            status = probe.solve(cbc)
            if pulp.LpStatus[status] == "Optimal":
                _SOLVER = pulp.PULP_CBC_CMD(msg=False, timeLimit=180)
                return _SOLVER
    except OSError:
        pass
    highs = pulp.HiGHS(msg=False, timeLimit=180)
    if not highs.available():
        raise SystemExit(
            "Aucun solveur utilisable. Installez highspy (pip install -r requirements.txt) "
            "ou un CBC compatible avec cette machine."
        )
    print("Solveur HiGHS : le CBC livré avec PuLP ne s'exécute pas sur cette machine.")
    _SOLVER = highs
    return _SOLVER


def supabase_headers(key: str) -> dict:
    return {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
    }


def fetch_all(url: str, key: str, table: str, select: str) -> list[dict]:
    rows: list[dict] = []
    start = 0
    page = 1000
    while True:
        response = requests.get(
            f"{url.rstrip('/')}/rest/v1/{table}",
            params={"select": select},
            headers={**supabase_headers(key), "Range": f"{start}-{start + page - 1}"},
            timeout=60,
        )
        response.raise_for_status()
        chunk = response.json()
        rows.extend(chunk)
        if len(chunk) < page:
            return rows
        start += page


def next_tuesday_0900() -> str:
    """Mardi suivant à 9 h, heure d'arrivée souhaitée."""
    now = datetime.now(ZoneInfo("Europe/Paris"))
    days = (1 - now.weekday()) % 7
    if days == 0 and now.hour >= 9:
        days = 7
    target = (now + timedelta(days=days)).replace(hour=9, minute=0, second=0, microsecond=0)
    return target.strftime("%Y%m%dT%H%M%S")


def pair_key(lat1: float, lon1: float, lat2: float, lon2: float) -> str:
    return f"{lat1:.5f},{lon1:.5f}|{lat2:.5f},{lon2:.5f}"


def haversine_minutes(lat1: float, lon1: float, lat2: float, lon2: float) -> int:
    radius = 6371.0
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    d_phi = math.radians(lat2 - lat1)
    d_lambda = math.radians(lon2 - lon1)
    a = math.sin(d_phi / 2) ** 2 + math.cos(phi1) * math.cos(phi2) * math.sin(d_lambda / 2) ** 2
    km = 2 * radius * math.asin(math.sqrt(a))
    # Vitesse mixte transports d'environ 30 km/h, quand Navitia ne répond pas.
    return max(1, round(km / 30 * 60))


def load_time_cache(when: str) -> dict:
    if not CACHE_PATH.exists():
        return {"datetime": when, "rules": TRIP_RULES, "pairs": {}}
    cache = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    cache["datetime"] = when
    cache["rules"] = TRIP_RULES
    cache.setdefault("pairs", {})
    return cache


def pair_is_fresh(cached) -> bool:
    if not isinstance(cached, dict):
        return False
    if cached.get("source") == "sncf":
        return True
    return cached.get("rules") == TRIP_RULES


def save_time_cache(cache: dict) -> None:
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CACHE_PATH.write_text(json.dumps(cache), encoding="utf-8")


def retry_after_seconds(header_text: str) -> int | None:
    for line in header_text.splitlines():
        name, _, value = line.partition(":")
        if name.strip().lower() == "retry-after":
            try:
                return max(0, int(value.strip()))
            except ValueError:
                return None
    return None


class RequestPace:
    """Espace les départs de requête, pour rester au débit demandé."""

    def __init__(self, per_second: int):
        self.gap = 1 / per_second
        self.lock = threading.Lock()
        self.ready = 0.0

    def wait(self) -> None:
        with self.lock:
            now = time.monotonic()
            start = max(now, self.ready)
            self.ready = start + self.gap
        delay = start - time.monotonic()
        if delay > 0:
            time.sleep(delay)


def section_hops(section: dict) -> int:
    return max(0, len(section.get("stop_date_times") or []) - 1)


def short_ride(section: dict) -> bool:
    """Métro d'une seule station, ou bus de moins de deux stations."""
    if section.get("type") != "public_transport":
        return False
    mode = ((section.get("display_informations") or {}).get("commercial_mode") or "").casefold()
    hops = section_hops(section)
    if "métro" in mode or "metro" in mode:
        return hops <= 1
    if mode == "bus":
        return hops < 2
    return False


def journey_span_minutes(journey: dict) -> int | None:
    """Durée du trajet, de son départ à son arrivée. L'attente ensuite n'est pas comptée."""
    start = journey.get("departure_date_time") or ""
    end = journey.get("arrival_date_time") or ""
    if len(start) >= 15 and len(end) >= 15:
        fmt = "%Y%m%dT%H%M%S"
        delta = datetime.strptime(end[:15], fmt) - datetime.strptime(start[:15], fmt)
        return max(1, round(delta.total_seconds() / 60))
    if journey.get("duration") is None:
        return None
    return max(1, round(journey["duration"] / 60))


def choose_journey(journeys: list[dict]) -> dict | None:
    """Trajet le plus court qui arrive pour 9 h. À durée égale, l'arrivée la plus tardive."""
    best = None
    best_minutes = None
    best_arrival = ""
    for journey in journeys:
        if any(short_ride(section) for section in journey.get("sections") or []):
            continue
        minutes = journey_span_minutes(journey)
        arrival = journey.get("arrival_date_time") or ""
        if minutes is None:
            continue
        if best is None or minutes < best_minutes or (minutes == best_minutes and arrival > best_arrival):
            best = journey
            best_minutes = minutes
            best_arrival = arrival
    return best


def arrival_clock(journey: dict) -> str | None:
    end = journey.get("arrival_date_time") or ""
    if len(end) < 15:
        return None
    return f"{end[9:11]}:{end[11:13]}:{end[13:15]}"


def navitia_journeys(api_key: str, origin, destination, when: str, pace: RequestPace | None, extra: dict) -> list[dict] | None:
    # curl, et non requests : le Python du Mac ne négocie pas le TLS exigé par Navitia.
    params = {
        "from": f"{origin[1]};{origin[0]}",
        "to": f"{destination[1]};{destination[0]}",
        "datetime": when,
        "datetime_represents": "arrival",
        "walking_speed": f"{WALKING_MPS:.3f}",
        "count": 10,
    }
    params.update(extra)
    query = urlencode(params)
    for attempt in range(4):
        if pace is not None:
            pace.wait()
        with tempfile.NamedTemporaryFile(mode="w+", encoding="utf-8") as header_file:
            completed = subprocess.run(
                [
                    "curl",
                    "-sS",
                    "--max-time",
                    "40",
                    "-D",
                    header_file.name,
                    "-H",
                    f"apiKey: {api_key}",
                    "-w",
                    "\n%{http_code}",
                    f"{NAVITIA_URL}?{query}",
                ],
                capture_output=True,
                text=True,
            )
            header_file.seek(0)
            headers = header_file.read()
        if completed.returncode != 0:
            time.sleep(2 ** attempt)
            continue
        body, _, status_text = completed.stdout.rpartition("\n")
        try:
            status = int(status_text.strip())
        except ValueError:
            return None
        if status in (401, 403):
            raise SystemExit("Clé API Île-de-France Mobilités refusée. Vérifiez IDFM_API_KEY dans algo/.env.")
        if status == 429:
            wait = retry_after_seconds(headers)
            if wait is None or wait > 120:
                raise SystemExit(
                    "Quota Navitia atteint. Le cache déjà obtenu est conservé ; "
                    f"relancez le calcul dans {wait or 'quelques heures'} secondes."
                )
            time.sleep(wait or (2 ** attempt))
            continue
        if status >= 500:
            time.sleep(2 ** attempt)
            continue
        if status != 200:
            return None
        try:
            return json.loads(body).get("journeys") or []
        except json.JSONDecodeError:
            return None
    return None


def navitia_choice(api_key: str, origin, destination, when: str, pace: RequestPace | None = None) -> tuple[int, str | None] | None:
    journeys = navitia_journeys(api_key, origin, destination, when, pace, {})
    chosen = choose_journey(journeys or []) if journeys is not None else None
    if chosen is None and journeys is not None:
        walking = navitia_journeys(api_key, origin, destination, when, pace, {"direct_path": "only", "count": 1})
        chosen = choose_journey(walking or [])
    if chosen is None:
        return None
    minutes = journey_span_minutes(chosen)
    if minutes is None:
        return None
    return minutes, arrival_clock(chosen)


def navitia_minutes(session: requests.Session, api_key: str, origin, destination, when: str, pace: RequestPace | None = None) -> int | None:
    del session
    found = navitia_choice(api_key, origin, destination, when, pace)
    if found is None:
        return None
    return found[0]


def load_car_cache() -> dict:
    if not CAR_CACHE_PATH.exists():
        return {}
    return json.loads(CAR_CACHE_PATH.read_text(encoding="utf-8"))


def save_car_cache(cache: dict) -> None:
    CAR_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CAR_CACHE_PATH.write_text(json.dumps(cache), encoding="utf-8")


def osrm_minutes(session: requests.Session, origin, destination) -> int | None:
    url = f"{OSRM_URL}/{origin[1]},{origin[0]};{destination[1]},{destination[0]}"
    for attempt in range(4):
        response = session.get(url, params={"overview": "false"}, timeout=40)
        if response.status_code == 429 or response.status_code >= 500:
            time.sleep(2 ** attempt)
            continue
        if response.status_code != 200:
            return None
        routes = response.json().get("routes") or []
        if not routes or routes[0].get("duration") is None:
            return None
        return max(1, round(routes[0]["duration"] / 60))
    return None


def transit_minutes(session, cache, api_key, origin, destination, when, counters) -> int:
    key = pair_key(origin[0], origin[1], destination[0], destination[1])
    cached = cache["pairs"].get(key)
    if pair_is_fresh(cached):
        return int(cached["minutes"])
    found = navitia_choice(api_key, origin, destination, when)
    source = "navitia"
    arrivee = None
    if found is None:
        minutes = haversine_minutes(origin[0], origin[1], destination[0], destination[1])
        source = "haversine"
        counters["fallbacks"] += 1
    else:
        minutes, arrivee = found
    cache["pairs"][key] = {"minutes": minutes, "source": source, "rules": TRIP_RULES, "arrivee": arrivee}
    counters["new_calls"] += 1
    save_time_cache(cache)
    time.sleep(0.1)
    return minutes


def car_minutes(session, cache, origin, destination) -> int | None:
    key = pair_key(origin[0], origin[1], destination[0], destination[1])
    cached = cache.get(key)
    if cached is not None:
        return int(cached)
    minutes = osrm_minutes(session, origin, destination)
    if minutes is None:
        return None
    cache[key] = minutes
    time.sleep(0.2)
    return minutes


def build_time_matrix(students: list[dict], stores: list[dict], api_key: str, when: str) -> tuple[list[list[int]], int]:
    cache = load_time_cache(when)
    cars = load_car_cache()
    session = requests.Session()
    matrix: list[list[int]] = []
    counters = {"fallbacks": 0, "new_calls": 0}
    done = 0
    real_students = sum(1 for student in students if not student.get("absent"))
    total = real_students * len(stores)
    print(f"  Navitia sur toutes les paires étudiant–magasin : {total} paires.")
    print(f"  Débit Navitia : {NAVITIA_PER_SECOND} requêtes par seconde. Le cache est réutilisé.")

    pending: dict[str, tuple[tuple[float, float], tuple[float, float], list[tuple[int, int]]]] = {}
    for student_index, student in enumerate(students):
        if student.get("absent"):
            matrix.append([0] * len(stores))
            continue
        origin = (student["latitude"], student["longitude"])
        row = [UNREACHABLE] * len(stores)
        for index in range(len(stores)):
            store = stores[index]
            destination = (store["latitude"], store["longitude"])
            key = pair_key(origin[0], origin[1], destination[0], destination[1])
            cached = cache["pairs"].get(key)
            if pair_is_fresh(cached):
                row[index] = int(cached["minutes"])
                done += 1
                continue
            slots = pending.setdefault(key, (origin, destination, []))
            slots[2].append((student_index, index))
        matrix.append(row)
        if student.get("a_voiture"):
            for index in range(len(stores)):
                if row[index] >= UNREACHABLE:
                    continue
                driving = car_minutes(session, cars, origin, (stores[index]["latitude"], stores[index]["longitude"]))
                if driving is not None:
                    row[index] = min(row[index], driving)

    if done:
        print(f"  {done}/{total} paires déjà en cache", flush=True)

    pace = RequestPace(NAVITIA_PER_SECOND)

    def fetch(key: str, origin, destination):
        found = navitia_choice(api_key, origin, destination, when, pace)
        source = "navitia"
        arrivee = None
        if found is None:
            minutes = haversine_minutes(origin[0], origin[1], destination[0], destination[1])
            source = "haversine"
        else:
            minutes, arrivee = found
        return key, minutes, source, arrivee

    if pending:
        with ThreadPoolExecutor(max_workers=NAVITIA_PER_SECOND * 3) as pool:
            futures = [
                pool.submit(fetch, key, origin, destination)
                for key, (origin, destination, _slots) in pending.items()
            ]
            try:
                for future in as_completed(futures):
                    key, minutes, source, arrivee = future.result()
                    _origin, _destination, slots = pending[key]
                    cache["pairs"][key] = {"minutes": minutes, "source": source, "rules": TRIP_RULES, "arrivee": arrivee}
                    counters["new_calls"] += 1
                    if source == "haversine":
                        counters["fallbacks"] += 1
                    for student_index, store_index in slots:
                        value = minutes
                        student = students[student_index]
                        if student.get("a_voiture"):
                            driving = car_minutes(
                                session,
                                cars,
                                (student["latitude"], student["longitude"]),
                                (stores[store_index]["latitude"], stores[store_index]["longitude"]),
                            )
                            if driving is not None:
                                value = min(value, driving)
                        current = matrix[student_index][store_index]
                        matrix[student_index][store_index] = value if current >= UNREACHABLE else min(current, value)
                        done += 1
                    save_time_cache(cache)
                    if done % 100 < len(slots) or done == total:
                        print(f"  {done}/{total} paires", flush=True)
            except SystemExit:
                save_time_cache(cache)
                pool.shutdown(wait=False, cancel_futures=True)
                raise

    save_time_cache(cache)
    save_car_cache(cars)
    return matrix, counters["fallbacks"]


def google_store_coordinates() -> dict[str, tuple[float, float]]:
    from openpyxl import load_workbook

    if not GOOGLE_ADDRESSES_PATH.exists():
        raise SystemExit(f"Fichier introuvable : {GOOGLE_ADDRESSES_PATH.name}")
    sheet = load_workbook(GOOGLE_ADDRESSES_PATH, data_only=True).active
    found = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name, latitude, longitude = row[0], row[5], row[6]
        if name and latitude is not None and longitude is not None:
            found[str(name)] = (float(latitude), float(longitude))
    return found


def build_google_matrix(students: list[dict], stores: list[dict]) -> tuple[list[list[int]], dict]:
    """Durées et heures d'arrivée du cache Google. Aucun appel Navitia."""
    if not GOOGLE_CACHE_PATH.exists():
        raise SystemExit(f"Cache Google introuvable : {GOOGLE_CACHE_PATH.name}")
    pairs = json.loads(GOOGLE_CACHE_PATH.read_text(encoding="utf-8")).get("pairs") or {}
    located = google_store_coordinates()
    missing_stores = [store["nom"] for store in stores if store["nom"] not in located]
    if missing_stores:
        raise SystemExit("Adresse Google manquante : " + ", ".join(missing_stores))

    matrix: list[list[int]] = []
    arrivals: dict[tuple[str, int], str] = {}
    missing = 0
    for student in students:
        if student.get("absent"):
            matrix.append([0] * len(stores))
            continue
        row = []
        for store in stores:
            latitude, longitude = located[store["nom"]]
            cached = pairs.get(pair_key(student["latitude"], student["longitude"], latitude, longitude))
            fresh = (
                isinstance(cached, dict)
                and cached.get("source") == "google"
                and cached.get("rules") == GOOGLE_RULES
            )
            if not fresh:
                missing += 1
                row.append(UNREACHABLE)
                continue
            row.append(int(cached["minutes"]))
            clock = cached.get("arrivee")
            if clock:
                arrivals[student["email"], store["id"]] = clock
        matrix.append(row)
    if missing:
        raise SystemExit(
            f"{missing} paires sans durée Google. Le cache Navitia n'est pas utilisé."
        )
    return matrix, arrivals


def check_problem_size(students: list[dict], stores: list[dict]) -> None:
    missing = [store["nom"] for store in stores if store.get("latitude") is None or store.get("longitude") is None]
    if missing:
        raise SystemExit("Magasins sans coordonnées : " + ", ".join(missing))

    present = {store["zone"] for store in stores}
    if set(ZONES) - present:
        raise SystemExit("Zones manquantes : " + ", ".join(sorted(set(ZONES) - present)))

    n = len(students)
    total_min = 0
    total_max = 0
    for zone in ZONES:
        group = [store for store in stores if store["zone"] == zone]
        ceiling = sum(store["capacite_max"] for store in group)
        low = ZONE_MIN
        high = ceiling
        if low > high:
            raise SystemExit(
                f"{zone} : les magasins accueillent au plus {ceiling} étudiants, "
                f"moins que le minimum de {ZONE_MIN}."
            )
        total_min += low
        total_max += high
        print(f"  {zone} : effectif visé {low}–{high} (plafond magasins {ceiling})")

    if not (total_min <= n <= total_max):
        raise SystemExit(
            f"{n} étudiants pour une fourchette globale de {total_min}–{total_max}. "
            f"Chaque zone doit recevoir au moins {ZONE_MIN} étudiants."
        )


def student_rank(student, store_id: int) -> int:
    if student.get("absent"):
        return 1
    return rank_of(student["voeux"], store_id)


def rank_of(voeux, store_id: int) -> int:
    try:
        ids = [int(item) for item in voeux]
    except (TypeError, ValueError):
        return OUTSIDE_RANK
    try:
        return ids.index(int(store_id)) + 1
    except ValueError:
        return OUTSIDE_RANK


def eligible_pairs(students, stores, matrix, max_minutes) -> set[tuple[int, int]]:
    """Paires autorisées. Un magasin sans étudiant dans le plafond peut rester vide."""
    allowed: set[tuple[int, int]] = set()
    for j in range(len(stores)):
        for i in range(len(students)):
            if matrix[i][j] >= UNREACHABLE:
                continue
            if max_minutes is not None and matrix[i][j] > max_minutes:
                continue
            allowed.add((i, j))
    return allowed


def build_model(students, stores, matrix, max_minutes, rank_caps):
    problem = pulp.LpProblem("affectation", pulp.LpMinimize)
    variables: dict[tuple[int, int], pulp.LpVariable] = {}
    allowed = eligible_pairs(students, stores, matrix, max_minutes)

    for i, j in sorted(allowed):
        variables[i, j] = pulp.LpVariable(f"x_{i}_{j}", cat="Binary")

    for i, student in enumerate(students):
        choices = [variables[i, j] for j in range(len(stores)) if (i, j) in variables]
        if not choices:
            return None
        problem += pulp.lpSum(choices) == 1, f"un_{i}"
        forced = FORCED_STORES.get(student["email"])
        if not forced:
            continue
        pinned = [
            variables[i, j]
            for j, store in enumerate(stores)
            if store["nom"] in forced and (i, j) in variables
        ]
        if not pinned:
            return None
        problem += pulp.lpSum(pinned) == 1, f"place_{i}"

    for j, store in enumerate(stores):
        hosted = [variables[i, j] for i in range(len(students)) if (i, j) in variables]
        if not hosted:
            continue
        opened = pulp.LpVariable(f"ouvert_{j}", cat="Binary")
        problem += pulp.lpSum(hosted) >= store["capacite_min"] * opened, f"cap_min_{j}"
        problem += pulp.lpSum(hosted) <= store["capacite_max"] * opened, f"cap_max_{j}"

    for zone in ZONES:
        indexes = [j for j, store in enumerate(stores) if store["zone"] == zone]
        index_set = set(indexes)
        reachable = {i for (i, j) in variables if j in index_set}
        if len(reachable) < ZONE_MIN:
            return None
        hosted = [variables[i, j] for j in indexes for i in range(len(students)) if (i, j) in variables]
        problem += pulp.lpSum(hosted) >= ZONE_MIN, f"zone_min_{zone}"

    for threshold, maximum in rank_caps:
        terms = []
        for i, student in enumerate(students):
            for j, store in enumerate(stores):
                if (i, j) in variables and student_rank(student, store["id"]) >= threshold:
                    terms.append(variables[i, j])
        problem += pulp.lpSum(terms) <= maximum, f"rang_{threshold}"

    return problem, variables


def optimal(problem):
    status = problem.solve(make_solver())
    if pulp.LpStatus[status] != "Optimal":
        return None
    return pulp.value(problem.objective)


def assignment_from(variables, students, stores) -> dict[str, int] | None:
    assignment: dict[str, int] = {}
    for i, student in enumerate(students):
        chosen = None
        for j, store in enumerate(stores):
            variable = variables.get((i, j))
            if variable is not None and variable.value() is not None and variable.value() > 0.5:
                chosen = store["id"]
                break
        if chosen is None:
            return None
        assignment[student["email"]] = chosen
    return assignment


def solve_sum(students, stores, matrix, max_minutes, rank_caps, tail=None):
    built = build_model(students, stores, matrix, max_minutes, rank_caps)
    if built is None:
        return None
    problem, variables = built
    if tail is not None:
        limit, allowed = tail
        overs = [pulp.LpVariable(f"queue_{i}", cat="Binary") for i in range(len(students))]
        problem += pulp.lpSum(overs) <= allowed
        for (i, j), variable in variables.items():
            if matrix[i][j] > limit:
                problem += variable <= overs[i]
    costs = []
    for (i, j), variable in variables.items():
        costs.append(matrix[i][j] * variable)
    problem += pulp.lpSum(costs)
    if optimal(problem) is None:
        return None
    return assignment_from(variables, students, stores)


def solve_worst(students, stores, matrix, max_minutes, rank_caps):
    built = build_model(students, stores, matrix, max_minutes, rank_caps)
    if built is None:
        return None
    problem, variables = built
    worst = pulp.LpVariable("pire", lowBound=0, cat="Integer")
    for (i, j), variable in variables.items():
        problem += worst >= matrix[i][j] * variable
    problem += worst
    value = optimal(problem)
    if value is None:
        return None
    limit = int(round(value))
    tightened = limit if max_minutes is None else min(max_minutes, limit)
    return solve_sum(students, stores, matrix, tightened, rank_caps)


def solve_p90(students, stores, matrix, max_minutes, rank_caps):
    built = build_model(students, stores, matrix, max_minutes, rank_caps)
    if built is None:
        return None
    problem, variables = built
    allowed = len(students) // 10
    big = max(matrix[i][j] for (i, j) in variables) + 1
    overs = [pulp.LpVariable(f"haut_{i}", cat="Binary") for i in range(len(students))]
    threshold = pulp.LpVariable("seuil", lowBound=0, cat="Integer")
    for (i, j), variable in variables.items():
        problem += matrix[i][j] * variable <= threshold + big * overs[i]
    problem += pulp.lpSum(overs) <= allowed
    problem += threshold
    value = optimal(problem)
    if value is None:
        return None
    return solve_sum(
        students,
        stores,
        matrix,
        max_minutes,
        rank_caps,
        tail=(int(round(value)), allowed),
    )


def solve_objective(kind, students, stores, matrix, max_minutes, rank_caps):
    if kind == "moyenne":
        return solve_sum(students, stores, matrix, max_minutes, rank_caps)
    if kind == "pire":
        return solve_worst(students, stores, matrix, max_minutes, rank_caps)
    if kind == "p90":
        return solve_p90(students, stores, matrix, max_minutes, rank_caps)
    raise ValueError(kind)


def lexico_rank_caps(students, stores, matrix, max_minutes):
    caps = []
    for threshold in RANK_STEPS:
        built = build_model(students, stores, matrix, max_minutes, caps)
        if built is None:
            return None
        problem, variables = built
        terms = []
        for i, student in enumerate(students):
            for j, store in enumerate(stores):
                if (i, j) in variables and student_rank(student, store["id"]) >= threshold:
                    terms.append(variables[i, j])
        problem += pulp.lpSum(terms)
        value = optimal(problem)
        if value is None:
            return None
        caps.append((threshold, int(round(value))))
    return caps


def solve_six(students, stores, matrix):
    caps = []
    if any(family == "voeux" for _number, _code, family, _kind in ALGORITHMS):
        print("Famille vœux : rangs lexicographiques.")
        caps = lexico_rank_caps(students, stores, matrix, None)
        if caps is None:
            raise SystemExit("Famille vœux infaisable avec les capacités et les zones.")
    results = {}
    for _number, code, family, kind in ALGORITHMS:
        print(f"  {code}")
        rank_caps = caps if family == "voeux" else []
        assignment = solve_objective(kind, students, stores, matrix, None, rank_caps)
        if assignment is None:
            raise SystemExit(f"{code} infaisable.")
        results[code] = assignment
    return results


def validate(assignment: dict[str, int], students: list[dict], stores: list[dict], matrix=None) -> dict[str, int]:
    by_id = {store["id"]: store for store in stores}
    if set(assignment) != {student["email"] for student in students}:
        raise SystemExit("L'affectation ne couvre pas exactement les étudiants.")

    per_store: dict[int, int] = {}
    per_zone: dict[str, int] = {zone: 0 for zone in ZONES}
    for store_id in assignment.values():
        store = by_id.get(store_id)
        if store is None:
            raise SystemExit(f"Magasin inconnu dans l'affectation : {store_id}")
        per_store[store_id] = per_store.get(store_id, 0) + 1
        per_zone[store["zone"]] += 1

    del matrix
    for store in stores:
        count = per_store.get(store["id"], 0)
        if count == 0:
            continue
        if not (store["capacite_min"] <= count <= store["capacite_max"]):
            raise SystemExit(
                f"{store['nom']} reçoit {count} étudiants, hors de "
                f"{store['capacite_min']}–{store['capacite_max']}."
            )
    for zone, count in per_zone.items():
        if count < ZONE_MIN:
            raise SystemExit(f"{zone} reçoit {count} étudiants, moins que le minimum de {ZONE_MIN}.")
    return per_zone


def time_lookup(students, stores, matrix) -> dict[tuple[str, int], int]:
    lookup = {}
    for i, student in enumerate(students):
        for j, store in enumerate(stores):
            lookup[student["email"], store["id"]] = matrix[i][j]
    return lookup


def print_report(title: str, assignment: dict[str, int], students, stores, matrix) -> int:
    zones = validate(assignment, students, stores, matrix)
    lookup = time_lookup(students, stores, matrix)
    times = [lookup[email, store_id] for email, store_id in assignment.items()]
    over = sum(1 for value in times if value > 45)
    print(f"{title}")
    for zone in ZONES:
        print(f"  {zone} : {zones[zone]} étudiants")
    print(f"  Temps moyen : {sum(times) / len(times):.1f} min, maximum : {max(times)} min")
    print(f"  Étudiants au-delà de 45 min : {over}")
    return over


def replace_results(url: str, key: str, rows: list[dict]) -> None:
    headers = {**supabase_headers(key), "Prefer": "return=minimal"}
    deleted = requests.delete(
        f"{url.rstrip('/')}/rest/v1/resultats_affectation",
        params={"id": "gt.0"},
        headers=headers,
        timeout=60,
    )
    if not deleted.ok:
        raise SystemExit(f"Suppression des anciens résultats refusée ({deleted.status_code}) : {deleted.text}")
    inserted = requests.post(
        f"{url.rstrip('/')}/rest/v1/resultats_affectation",
        headers=headers,
        json=rows,
        timeout=60,
    )
    if not inserted.ok:
        raise SystemExit(f"Insertion des résultats refusée ({inserted.status_code}) : {inserted.text}")


def publish(url: str, key: str) -> None:
    response = requests.patch(
        f"{url.rstrip('/')}/rest/v1/app_config",
        params={"id": "eq.1"},
        headers={**supabase_headers(key), "Prefer": "return=minimal"},
        json={"phase": "vote", "updated_at": datetime.now(ZoneInfo("Europe/Paris")).isoformat()},
        timeout=60,
    )
    if not response.ok:
        raise SystemExit(f"Passage en phase vote refusé ({response.status_code}) : {response.text}")
    print("Phase publiée : vote.")


def normalize_students(rows: list[dict]) -> list[dict]:
    students = []
    for row in rows:
        students.append(
            {
                "email": str(row["email"]).strip().lower(),
                "latitude": float(row["latitude"]),
                "longitude": float(row["longitude"]),
                "voeux": json.loads(row["voeux"]) if isinstance(row["voeux"], str) else row["voeux"],
            }
        )
    return students


def normalize_stores(rows: list[dict]) -> list[dict]:
    stores = []
    for row in rows:
        stores.append(
            {
                "id": int(row["id"]),
                "nom": row["nom"],
                "zone": row["zone"],
                "latitude": None if row["latitude"] is None else float(row["latitude"]),
                "longitude": None if row["longitude"] is None else float(row["longitude"]),
                "capacite_min": int(row["capacite_min"]),
                "capacite_max": int(row["capacite_max"]),
            }
        )
    return stores


def build_rows(students, results, lookup, arrivals=None) -> list[dict]:
    rows = []
    for student in students:
        if student.get("absent"):
            continue
        email = student["email"]
        for _number, code, _family, _kind in ALGORITHMS:
            store_id = results[code][email]
            clock = None if arrivals is None else arrivals.get((email, store_id))
            rows.append({
                "email_etudiant": email,
                "algo": code,
                "magasin_id": store_id,
                "temps_min": lookup[email, store_id],
                "rang_voeu": student_rank(student, store_id),
                "arrivee": clock,
            })
    return rows


def histogram_upper(values) -> int:
    peak = max(values) if values else 0
    return max(5, (int(peak) // 5 + 1) * 5)


def time_histogram(values, upper: int) -> list[tuple[int, int, int]]:
    bins = []
    start = 0
    while start < upper:
        end = start + 5
        count = sum(1 for value in values if start <= value < end)
        bins.append((start, end, count))
        start = end
    return bins


def rank_profile(assignment, students) -> tuple:
    voeux = {student["email"]: student["voeux"] for student in students}
    return tuple(sorted((rank_of(voeux[email], store_id) for email, store_id in assignment.items()), reverse=True))


def travel_times(assignment, lookup) -> list[int]:
    return [lookup[email, store_id] for email, store_id in assignment.items()]


def percentile_threshold(values) -> int:
    ordered = sorted(values)
    allowed = len(values) // 10
    return ordered[len(values) - allowed - 1]


def check_family(students, stores, matrix, results) -> None:
    lookup = time_lookup(students, stores, matrix)
    wish_profiles = []
    address_profiles = []
    all_times = []
    by_code = {}
    for _number, code, family, _kind in ALGORITHMS:
        validate(results[code], students, stores)
        print_report(code, results[code], students, stores, matrix)
        times = travel_times(results[code], lookup)
        by_code[code] = times
        all_times.extend(times)
        profile = rank_profile(results[code], students)
        if family == "voeux":
            wish_profiles.append(profile)
        else:
            address_profiles.append(profile)
    if wish_profiles and len(set(wish_profiles)) != 1:
        raise SystemExit("Les trois algorithmes de vœux n'ont pas le même profil de rangs.")
    if wish_profiles and wish_profiles[0] > min(address_profiles):
        raise SystemExit("La famille vœux devrait avoir un profil de rangs au moins aussi bon.")
    families = []
    for _number, _code, family, _kind in ALGORITHMS:
        if family not in families:
            families.append(family)
    for family in families:
        if not all(f"{family}_{kind}" in by_code for kind in ("moyenne", "pire", "p90")):
            continue
        mean = by_code[f"{family}_moyenne"]
        worst = by_code[f"{family}_pire"]
        middle = by_code[f"{family}_p90"]
        if sum(mean) > sum(worst) or sum(mean) > sum(middle):
            raise SystemExit(f"La moyenne de {family} devrait avoir la plus petite somme.")
        if max(worst) > max(mean):
            raise SystemExit(f"Le pire cas de {family} devrait avoir le plus petit maximum.")
        if percentile_threshold(middle) > percentile_threshold(mean):
            raise SystemExit(f"Le 90e percentile de {family} devrait être le plus bas.")
    upper = histogram_upper(all_times)
    for times in by_code.values():
        bins = time_histogram(times, upper)
        if sum(count for _start, _end, count in bins) != len(times):
            raise SystemExit("Les tranches de 5 minutes ne couvrent pas tous les étudiants.")
        if any(end - start != 5 for start, end, _count in bins):
            raise SystemExit("Une tranche ne fait pas 5 minutes.")


def self_test() -> None:
    stores = []
    store_id = 1
    for zone in ZONES:
        for _ in range(8):
            stores.append(
                {
                    "id": store_id,
                    "nom": f"Magasin {store_id}",
                    "zone": zone,
                    "latitude": 48.8,
                    "longitude": 2.3,
                    "capacite_min": 2,
                    "capacite_max": 4,
                }
            )
            store_id += 1

    students = []
    wish_ids = [store["id"] for store in stores[:10]]
    for index in range(120):
        students.append(
            {
                "email": f"etudiant{index}@exemple.fr",
                "latitude": 48.85,
                "longitude": 2.35,
                "voeux": wish_ids,
            }
        )

    short = [[15 + (j % 5) for j in range(len(stores))] for _ in students]
    print("Auto-test, trajets courts.")
    check_problem_size(students, stores)
    short_results = solve_six(students, stores, short)
    check_family(students, stores, short, short_results)

    far = []
    for row in short:
        far.append([90 if stores[j]["zone"] == "DIF04" else value for j, value in enumerate(row)])
    print("Auto-test, zone DIF04 au-delà de 45 min.")
    far_results = solve_six(students, stores, far)
    check_family(students, stores, far, far_results)
    lookup = time_lookup(students, stores, far)
    over = sum(1 for value in travel_times(far_results["adresses_pire"], lookup) if value > 45)
    if over < ZONE_MIN:
        raise SystemExit(f"DIF04 devrait compter au moins {ZONE_MIN} trajets au-delà de 45 min.")
    print("Auto-test réussi.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Calcule et publie les affectations par adresse.")
    parser.add_argument("--publier", action="store_true", help="Passe la phase à « vote » après l'écriture")
    parser.add_argument("--self-test", action="store_true", help="Vérifie le solveur sur un jeu synthétique")
    parser.add_argument("--avec-voiture", action="store_true", help="Tient compte des voitures déclarées")
    parser.add_argument("--absents", type=int, default=0, help="Étudiants non inscrits, temps 0 et vœu n°1 partout")
    args = parser.parse_args()

    if args.self_test:
        self_test()
        return

    load_dotenv(ROOT / ".env")
    url = os.environ.get("SUPABASE_URL", "").strip()
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
    if not url or not key:
        raise SystemExit("Renseignez SUPABASE_URL et SUPABASE_SERVICE_ROLE_KEY dans algo/.env.")

    students = normalize_students(
        fetch_all(url, key, "etudiants_voeux", "email,latitude,longitude,voeux")
    )
    cars = {
        str(row["email"]).strip().lower(): bool(row["a_voiture"])
        for row in fetch_all(url, key, "mobilite", "email,a_voiture")
    }
    for student in students:
        student["absent"] = False
        student["a_voiture"] = cars.get(student["email"], False) if args.avec_voiture else False
    if args.absents:
        for index in range(args.absents):
            students.append(
                {
                    "email": f"absent-{index + 1}@interne",
                    "latitude": 0.0,
                    "longitude": 0.0,
                    "voeux": [],
                    "a_voiture": False,
                    "absent": True,
                }
            )
        print(f"{args.absents} étudiants non inscrits : temps 0, chaque magasin est leur vœu n°1.")
    if args.avec_voiture:
        print(f"Étudiants avec voiture : {sum(1 for student in students if student['a_voiture'])}.")
    else:
        print("Voitures ignorées pour ce calcul.")
    raw_stores = fetch_all(
        url,
        key,
        "magasins",
        "id,nom,zone,code_postal,latitude,longitude,capacite_min,capacite_max",
    )
    idf_stores = []
    for row in raw_stores:
        if str(row["code_postal"])[:2] in IDF_DEPTS:
            idf_stores.append(row)
            continue
        print(f"  Magasin hors Île-de-France retiré : {row['nom']}")
    stores = normalize_stores(idf_stores)
    if not students:
        raise SystemExit("Aucun vœu en base.")
    if not stores:
        raise SystemExit("Aucun magasin en base.")

    print(f"{len(students)} étudiants, {len(stores)} magasins.")
    known = {student["email"] for student in students}
    for email, names in FORCED_STORES.items():
        if email not in known:
            raise SystemExit(f"Affectation imposée impossible : {email} n'est pas inscrit.")
        print(f"  {email} est placé à {' ou '.join(names)}.")
    check_problem_size(students, stores)

    when = next_tuesday_0900()
    print(f"Horaires de référence : arrivée le mardi {when}.")
    print("  Durées Google en transports. Le cache Navitia n'est pas lu.")
    matrix, arrivals = build_google_matrix(students, stores)
    print(f"  {sum(len(row) for row in matrix)} durées Google.")

    results = solve_six(students, stores, matrix)
    lookup = time_lookup(students, stores, matrix)
    for _number, code, _family, _kind in ALGORITHMS:
        print_report(code, results[code], students, stores, matrix)
    rows = build_rows(students, results, lookup, arrivals)
    replace_results(url, key, rows)
    print(f"{len(rows)} résultats écrits dans resultats_affectation.")
    if args.publier:
        publish(url, key)
    else:
        print("Phase inchangée.")


if __name__ == "__main__":
    try:
        main()
    except requests.RequestException as exc:
        raise SystemExit(f"Erreur réseau : {exc}") from exc
