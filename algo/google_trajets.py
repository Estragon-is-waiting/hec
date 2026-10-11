#!/usr/bin/env python3
"""Deuxième calcul des durées, via l'API Google Routes en transports en commun.

Le cache Navitia n'est pas modifié. Les durées Google vont dans cache/temps_google.json.
L'arrivée visée est le mardi à 9 h, comme pour Navitia. La durée est celle du trajet.
Un métro d'une seule station et un bus de moins de deux stations sont écartés.
Google choisit sa propre vitesse de marche.

  python google_trajets.py --essai 2
  python google_trajets.py --tous
"""

from __future__ import annotations

import argparse
import json
import math
import os
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from zoneinfo import ZoneInfo

import requests
from dotenv import load_dotenv

from run_affectation import (
    IDF_DEPTS,
    ROOT,
    fetch_all,
    next_tuesday_0900,
    pair_key,
)

CACHE_PATH = ROOT / "cache" / "temps_google.json"
NAVITIA_CACHE_PATH = ROOT / "cache" / "temps_trajet.json"
ROUTES_URL = "https://routes.googleapis.com/directions/v2:computeRoutes"
FIELD_MASK = ",".join(
    (
        "routes.duration",
        "routes.legs.steps.travelMode",
        "routes.legs.steps.transitDetails.stopDetails.arrivalTime",
        "routes.legs.steps.transitDetails.stopDetails.departureTime",
        "routes.legs.steps.transitDetails.stopCount",
        "routes.legs.steps.transitDetails.transitLine.vehicle.type",
    )
)
RULES = "arrivee-9h-google-transit"
PARIS = ZoneInfo("Europe/Paris")
ADDRESSES_PATH = ROOT.parent / "adresses_google.xlsx"
SAME_PLACE_METERS = 200
WORKERS = 24
_THREAD = threading.local()


def load_cache(when: str) -> dict:
    if not CACHE_PATH.exists():
        return {"datetime": when, "rules": RULES, "pairs": {}}
    cache = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    cache["datetime"] = when
    cache["rules"] = RULES
    cache.setdefault("pairs", {})
    return cache


def save_cache(cache: dict) -> None:
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CACHE_PATH.write_text(json.dumps(cache), encoding="utf-8")


def fresh(cached) -> bool:
    return isinstance(cached, dict) and cached.get("rules") == RULES and cached.get("source") == "google"


def arrival_stamp(when: str) -> str:
    moment = datetime.strptime(when, "%Y%m%dT%H%M%S").replace(tzinfo=PARIS)
    return moment.isoformat()


def duration_seconds(value) -> int | None:
    if isinstance(value, str) and value.endswith("s") and value[:-1].replace(".", "", 1).isdigit():
        return int(float(value[:-1]))
    return None


def short_step(step: dict) -> bool:
    """Métro d'une seule station, ou bus de moins de deux stations."""
    if step.get("travelMode") != "TRANSIT":
        return False
    details = step.get("transitDetails") or {}
    vehicle = ((details.get("transitLine") or {}).get("vehicle") or {}).get("type") or ""
    stops = int(details.get("stopCount") or 0)
    if vehicle in {"SUBWAY", "METRO_RAIL"}:
        return stops <= 1
    if vehicle == "BUS":
        return stops < 2
    return False


def route_span(route: dict) -> tuple[int, str | None, datetime | None] | None:
    seconds = duration_seconds(route.get("duration"))
    if seconds is None:
        return None
    arrivals: list[datetime] = []
    for leg in route.get("legs") or []:
        for step in leg.get("steps") or []:
            stop = ((step.get("transitDetails") or {}).get("stopDetails") or {})
            raw = stop.get("arrivalTime")
            if isinstance(raw, str):
                arrivals.append(datetime.fromisoformat(raw.replace("Z", "+00:00")))
    latest = max(arrivals) if arrivals else None
    clock = latest.astimezone(PARIS).strftime("%H:%M:%S") if latest else None
    return max(1, round(seconds / 60)), clock, latest


def choose_route(routes: list[dict]) -> tuple[int, str | None] | None:
    """Trajet acceptable le plus court. À durée égale, l'arrivée la plus tardive."""
    best = None
    best_minutes = None
    best_arrival = None
    for route in routes:
        steps = [step for leg in route.get("legs") or [] for step in leg.get("steps") or []]
        if any(short_step(step) for step in steps):
            continue
        parsed = route_span(route)
        if parsed is None:
            continue
        minutes, clock, arrival = parsed
        if best is None or minutes < best_minutes or (minutes == best_minutes and arrival and (best_arrival is None or arrival > best_arrival)):
            best = (minutes, clock)
            best_minutes = minutes
            best_arrival = arrival
    return best


def thread_session() -> requests.Session:
    session = getattr(_THREAD, "session", None)
    if session is None:
        session = requests.Session()
        _THREAD.session = session
    return session


def meters_between(origin, destination) -> float:
    radius = 6371000
    phi1, phi2 = math.radians(origin[0]), math.radians(destination[0])
    d_phi = math.radians(destination[0] - origin[0])
    d_lambda = math.radians(destination[1] - origin[1])
    arc = math.sin(d_phi / 2) ** 2 + math.cos(phi1) * math.cos(phi2) * math.sin(d_lambda / 2) ** 2
    return 2 * radius * math.asin(math.sqrt(arc))


def compute_routes(session: requests.Session, api_key: str, origin, destination, arrival: str, mode: str) -> list[dict]:
    for attempt in range(6):
        response = session.post(
            ROUTES_URL,
            headers={
                "Content-Type": "application/json",
                "X-Goog-Api-Key": api_key,
                "X-Goog-FieldMask": FIELD_MASK,
            },
            json={
                "origin": {"location": {"latLng": {"latitude": origin[0], "longitude": origin[1]}}},
                "destination": {"location": {"latLng": {"latitude": destination[0], "longitude": destination[1]}}},
                "travelMode": mode,
                "arrivalTime": arrival,
                "computeAlternativeRoutes": mode == "TRANSIT",
                "languageCode": "fr",
                "regionCode": "FR",
            },
            timeout=40,
        )
        try:
            payload = response.json()
        except json.JSONDecodeError:
            payload = {}
        if response.status_code == 200:
            return payload.get("routes") or []
        error = payload.get("error") or {}
        status = error.get("status") or response.status_code
        message = error.get("message") or "réponse vide"
        if response.status_code == 429 or status == "RESOURCE_EXHAUSTED":
            time.sleep(min(30, 2 ** attempt))
            continue
        raise RuntimeError(f"Google Routes a refusé l'appel ({status}). {message}")
    raise RuntimeError("Quota Google atteint. Le cache déjà obtenu est conservé.")


def google_choice(session, api_key: str, origin, destination, arrival: str) -> tuple[int, str | None] | None:
    chosen = choose_route(compute_routes(session, api_key, origin, destination, arrival, "TRANSIT"))
    if chosen is not None:
        return chosen
    walking = compute_routes(session, api_key, origin, destination, arrival, "WALK")
    if not walking:
        return None
    parsed = route_span(walking[0])
    if parsed is None:
        return None
    minutes, clock, _arrival = parsed
    return minutes, clock or "09:00:00"


def navitia_minutes(origin, destination) -> int | None:
    if not NAVITIA_CACHE_PATH.exists():
        return None
    cache = json.loads(NAVITIA_CACHE_PATH.read_text(encoding="utf-8"))
    cached = cache.get("pairs", {}).get(pair_key(origin[0], origin[1], destination[0], destination[1]))
    if isinstance(cached, dict) and cached.get("source") == "navitia":
        return int(cached["minutes"])
    return None


def google_addresses() -> dict[str, tuple[float, float]]:
    from openpyxl import load_workbook

    if not ADDRESSES_PATH.exists():
        raise SystemExit(f"Fichier introuvable : {ADDRESSES_PATH.name}")
    sheet = load_workbook(ADDRESSES_PATH, data_only=True).active
    found = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name, latitude, longitude = row[0], row[5], row[6]
        if name and latitude is not None and longitude is not None:
            found[str(name)] = (float(latitude), float(longitude))
    return found


def run_corrected(students, stores, api_key: str, when: str) -> None:
    """Tous les magasins, aux coordonnées Google. Le cache déjà juste est conservé."""
    located = google_addresses()
    missing = [store["nom"] for store in stores if store["nom"] not in located]
    if missing:
        raise SystemExit("Adresse Google manquante : " + ", ".join(missing))

    arrival = arrival_stamp(when)
    cache = load_cache(when)
    pending = []
    reused = 0
    kept = []
    moved = []
    for store in stores:
        previous = (float(store["latitude"]), float(store["longitude"]))
        current = located[store["nom"]]
        store["destination"] = current
        same = meters_between(previous, current) <= SAME_PLACE_METERS
        (kept if same else moved).append(store["nom"])
        for student in students:
            origin = (float(student["latitude"]), float(student["longitude"]))
            key = pair_key(origin[0], origin[1], current[0], current[1])
            if fresh(cache["pairs"].get(key)):
                reused += 1
                continue
            if same:
                former = pair_key(origin[0], origin[1], previous[0], previous[1])
                cached = cache["pairs"].get(former)
                if fresh(cached):
                    cache["pairs"][key] = cached
                    reused += 1
                    continue
            pending.append((key, origin, current, store["nom"]))

    save_cache(cache)
    print(f"{len(students)} étudiants, {len(stores)} magasins.")
    print(f"Adresses inchangées, cache conservé : {', '.join(kept)}.")
    print(f"Adresses corrigées, trajets à refaire : {', '.join(moved)}.")
    print(f"{reused} durées reprises, {len(pending)} appels en parallèle ({WORKERS} à la fois).")
    if not pending:
        print("Rien à appeler.")
        print("Le calcul d'affectation n'est pas relancé.")
        return

    lock = threading.Lock()
    called = 0
    failures = 0

    def fetch(item):
        key, origin, destination, name = item
        try:
            found = google_choice(thread_session(), api_key, origin, destination, arrival)
        except (requests.RequestException, RuntimeError) as exc:
            return key, name, None, str(exc)
        return key, name, found, None

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = [pool.submit(fetch, item) for item in pending]
        for future in as_completed(futures):
            key, name, found, error = future.result()
            with lock:
                called += 1
                if error or found is None:
                    failures += 1
                    if error:
                        print(f"  {name}: {error}", flush=True)
                else:
                    minutes, clock = found
                    cache["pairs"][key] = {
                        "minutes": minutes,
                        "arrivee": clock,
                        "source": "google",
                        "rules": RULES,
                    }
                if called % 40 == 0:
                    save_cache(cache)
                    print(f"  {called}/{len(pending)}", flush=True)
    save_cache(cache)
    print(f"{reused + called - failures} durées, {called} appels, {failures} sans itinéraire.")
    print("Le calcul d'affectation n'est pas relancé.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Durées Google en transports, à côté du calcul Navitia.")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--essai", type=int, metavar="N", help="N étudiants, tous les magasins d'Île-de-France")
    group.add_argument("--tous", action="store_true", help="Tous les étudiants inscrits")
    group.add_argument("--corriges", action="store_true", help="Tous les magasins, adresses Google, appels parallèles")
    args = parser.parse_args()

    load_dotenv(ROOT / ".env")
    url = os.environ.get("SUPABASE_URL", "").strip()
    service_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
    api_key = os.environ.get("GOOGLE_MAPS_API_KEY", "").strip()
    if not url or not service_key:
        raise SystemExit("Renseignez SUPABASE_URL et SUPABASE_SERVICE_ROLE_KEY dans algo/.env.")
    if not api_key:
        raise SystemExit("Renseignez GOOGLE_MAPS_API_KEY dans algo/.env. Aucun appel n'a été envoyé.")

    students = fetch_all(url, service_key, "etudiants_voeux", "email,nom,prenom,latitude,longitude")
    students.sort(key=lambda row: (str(row["nom"]).casefold(), str(row["prenom"]).casefold(), str(row["email"]).casefold()))
    if args.essai is not None:
        if args.essai < 1:
            raise SystemExit("L'essai demande au moins un étudiant.")
        students = students[: args.essai]
    stores = [
        row
        for row in fetch_all(url, service_key, "magasins", "id,nom,code_postal,latitude,longitude")
        if row.get("latitude") is not None and (args.corriges or str(row["code_postal"])[:2] in IDF_DEPTS)
    ]
    stores.sort(key=lambda row: str(row["nom"]).casefold())

    when = next_tuesday_0900()
    if args.corriges:
        run_corrected(students, stores, api_key, when)
        return
    arrival = arrival_stamp(when)
    cache = load_cache(when)
    session = requests.Session()
    total = len(students) * len(stores)
    print(f"{len(students)} étudiants, {len(stores)} magasins d'Île-de-France, {total} appels au plus.")
    print(f"Arrivée visée : mardi {when}.")

    done = 0
    called = 0
    missing = 0
    for student in students:
        origin = (float(student["latitude"]), float(student["longitude"]))
        print(f"\n{student['prenom']} {student['nom']}")
        for store in stores:
            destination = (float(store["latitude"]), float(store["longitude"]))
            key = pair_key(origin[0], origin[1], destination[0], destination[1])
            cached = cache["pairs"].get(key)
            if fresh(cached):
                minutes = int(cached["minutes"])
                clock = cached.get("arrivee") or "—"
            else:
                found = google_choice(session, api_key, origin, destination, arrival)
                called += 1
                if found is None:
                    missing += 1
                    print(f"  {store['nom']}: aucun itinéraire")
                    time.sleep(0.1)
                    continue
                minutes, clock = found
                cache["pairs"][key] = {
                    "minutes": minutes,
                    "arrivee": clock,
                    "source": "google",
                    "rules": RULES,
                }
                save_cache(cache)
                time.sleep(0.1)
            done += 1
            reference = navitia_minutes(origin, destination)
            compared = f"  Navitia {reference} min" if reference is not None else ""
            print(f"  {store['nom']}: Google {minutes} min, arrivée {clock or '—'}{compared}")

    print(f"\n{done} durées, {called} appels Google, {missing} sans itinéraire.")
    print("Le calcul d'affectation n'est pas relancé.")


if __name__ == "__main__":
    try:
        main()
    except requests.RequestException as exc:
        raise SystemExit(f"Erreur réseau : {exc}") from exc
